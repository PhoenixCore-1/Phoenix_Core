"""Customer Master file parsing and validation."""

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Any


@dataclass
class ParsedCustomerRow:
    row_number: int
    values: list[Any]
    external_id: str
    display_name: str
    validation_status: str
    validation_errors: list[str]
    validation_warnings: list[str]


@dataclass
class CustomerImportParseResult:
    headers: list[str]
    rows: list[ParsedCustomerRow]

    @property
    def total_rows(self) -> int:
        return len(self.rows)

    @property
    def valid_rows(self) -> int:
        return sum(
            row.validation_status == "VALID"
            for row in self.rows
        )

    @property
    def invalid_rows(self) -> int:
        return self.total_rows - self.valid_rows

    @property
    def preview_rows(self) -> list[ParsedCustomerRow]:
        return self.rows[:25]


def _read_csv(path: Path) -> list[list[Any]]:
    with path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as handle:
        return [list(row) for row in csv.reader(handle)]


def _read_excel(path: Path) -> list[list[Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    try:
        if not workbook.sheetnames:
            raise ValueError(
                "The file does not contain a worksheet."
            )

        sheet_name = (
            "Customers"
            if "Customers" in workbook.sheetnames
            else workbook.sheetnames[0]
        )

        worksheet = workbook[sheet_name]

        return [
            list(row)
            for row in worksheet.iter_rows(
                values_only=True,
            )
        ]

    finally:
        workbook.close()


def _read_matrix(path: Path) -> list[list[Any]]:
    suffix = path.suffix.lower()

    if suffix in {".xlsx", ".xlsm"}:
        return _read_excel(path)

    if suffix == ".csv":
        return _read_csv(path)

    raise ValueError(
        "Unsupported Customer Master file type. "
        "Only .xlsx, .xlsm, and .csv files are supported."
    )


def parse_customer_master_file(
    path: str | Path,
) -> CustomerImportParseResult:
    file_path = Path(path)

    matrix = _read_matrix(file_path)

    if not matrix:
        raise ValueError(
            "The selected file contains no rows."
        )

    headers = [
        str(value or "").strip()
        for value in matrix[0]
    ]

    customer_id_index = (
        headers.index("Customer ID")
        if "Customer ID" in headers
        else -1
    )

    customer_name_index = (
        headers.index("Customer name")
        if "Customer name" in headers
        else -1
    )

    missing_columns = []

    if customer_id_index == -1:
        missing_columns.append("Customer ID")

    if customer_name_index == -1:
        missing_columns.append("Customer name")

    if missing_columns:
        raise ValueError(
            "Missing required column(s): "
            + ", ".join(missing_columns)
        )

    source_rows = matrix[1:]

    customer_ids: dict[str, list[int]] = {}

    for index, values in enumerate(source_rows):
        raw_id = (
            values[customer_id_index]
            if customer_id_index < len(values)
            else None
        )

        customer_id = str(raw_id or "").strip()

        if customer_id:
            customer_ids.setdefault(
                customer_id,
                [],
            ).append(index + 2)

    rows = []

    for index, values in enumerate(source_rows):
        row_number = index + 2

        errors: list[str] = []
        warnings: list[str] = []

        raw_id = (
            values[customer_id_index]
            if customer_id_index < len(values)
            else None
        )

        raw_name = (
            values[customer_name_index]
            if customer_name_index < len(values)
            else None
        )

        customer_id = str(raw_id or "").strip()
        customer_name = str(raw_name or "").strip()

        if not customer_id:
            errors.append(
                "Customer ID is required."
            )
        else:
            duplicate_rows = customer_ids.get(
                customer_id,
                [],
            )

            if len(duplicate_rows) > 1:
                other_rows = [
                    value
                    for value in duplicate_rows
                    if value != row_number
                ]

                errors.append(
                    f'Duplicate Customer ID "{customer_id}" '
                    f"also appears on row {other_rows[0]}."
                )

        if not customer_name:
            errors.append(
                "Customer name is required."
            )

        rows.append(
            ParsedCustomerRow(
                row_number=row_number,
                values=values,
                external_id=customer_id,
                display_name=customer_name,
                validation_status=(
                    "INVALID"
                    if errors
                    else "VALID"
                ),
                validation_errors=errors,
                validation_warnings=warnings,
            )
        )

    return CustomerImportParseResult(
        headers=headers,
        rows=rows,
    )